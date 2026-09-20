"""Offline, re-runnable build for the federal court statistics supplement (round 3, generic view contract).

Inputs (never modified): seeds.jsonl, receipts/packet_*.jsonl, raw/**.
Outputs: tables.jsonl, cjra_rows.jsonl, gaps.json, validation.json, integration_spec.json.

Every figure carried here is a publisher-reported count for the stated period or as-of date. Nothing is
computed into a rate; percent columns that appear in a preview are the publisher's own cells, shown as printed.
No network access. Third-party packages used: openpyxl (XLSX preview), pypdf (PDF text); both were already installed.
"""
from __future__ import annotations

import calendar
import hashlib
import json
import re
import sys
import time
from collections import Counter, defaultdict
from datetime import date, datetime, time as dtime, timezone
from pathlib import Path

import openpyxl
from pypdf import PdfReader

HERE = Path(__file__).resolve().parent
RAW = (HERE / 'raw').resolve()
CACHE = HERE / '_work' / 'cjra_cache'
PARSER_VERSION = 'r3.2'
MAX_SHEETS, MAX_DATA_ROWS, MAX_HEADER_ROWS, MAX_COLS, MAX_CELL = 3, 40, 8, 40, 400

SERIES = {
    'Statistical Tables For The Federal Judiciary': ('stfj', 'Statistical Tables'),
    'Judicial Business': ('jb', 'Judicial Business'),
    'Federal Judicial Caseload Statistics': ('fjcs', 'Caseload Statistics'),
    'Judicial Facts and Figures': ('jff', 'Facts & Figures'),
    'Federal Court Management Statistics': ('fcms', 'FCMS'),
    'Civil Justice Reform Act (CJRA)': ('cjra', 'CJRA'),
    'Authorized Judgeships': ('judgeships', 'Judgeships'),
    'Judicial Vacancies': ('judgeships', 'Judgeships'),
}
TOPICS = {
    'civil_caseload_quarterly': 'Civil caseload', 'civil_caseload_annual': 'Civil caseload',
    'civil_trend_multiyear': 'Civil trends (multi-year)', 'trials_time_to_disposition': 'Trials and time to disposition',
    'judges_judgeships': 'Judges and judgeships', 'mdl_statistics': 'Multidistrict litigation',
    'district_court_management_profiles': 'District profiles', 'methodology': 'Methodology',
    'cjra_summary': 'CJRA summary', 'cjra_per_judge': 'CJRA per-judge',
}
MONTHS = {m: i for i, m in enumerate(calendar.month_name) if m}
DATE_RE = re.compile(r'(January|February|March|April|May|June|July|August|September|October|November|December)\s+(\d{1,2}),\s*(\d{4})', re.I)
LEAD_RE = re.compile(r'(during the|for the|as of|fiscal years?|calendar years?|\d+-month periods?|twelve-month periods?)', re.I)
CJRA_PER_JUDGE = {
    'CJRA 7': ('cjra7', 'All Cases', 'civil cases pending more than three years', 'civil'),
    'CJRA 8': ('cjra8', 'Motions', 'motions pending more than six months', 'motions'),
    'CJRA 9': ('cjra9', 'Bench Trials', 'bench trials submitted more than six months', 'bench'),
}
CIRCUITS = {'DC': 'District of Columbia', '1st': 'First', '2nd': 'Second', '3rd': 'Third', '4th': 'Fourth', '5th': 'Fifth', '6th': 'Sixth',
            '7th': 'Seventh', '8th': 'Eighth', '9th': 'Ninth', '10th': 'Tenth', '11th': 'Eleventh'}


def sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with path.open('rb') as fh:
        for chunk in iter(lambda: fh.read(1 << 20), b''):
            h.update(chunk)
    return h.hexdigest()


def read_jsonl(path: Path):
    return [json.loads(l) for l in path.read_text(encoding='utf-8').splitlines() if l.strip()]


def text_clean(value) -> str:
    return re.sub(r'\s+', ' ', str(value).replace('_x000D_', ' ')).strip()


def cell(value):
    if value is None or isinstance(value, bool):
        return value
    if isinstance(value, int):
        return value
    if isinstance(value, float):
        if value != value or value in (float('inf'), float('-inf')):
            return None
        return int(value) if value.is_integer() and abs(value) < 1e15 else round(value, 6)
    if isinstance(value, (datetime, date, dtime)):
        return value.isoformat()
    return text_clean(value)[:MAX_CELL] or None


YEAR_CONT_RE = re.compile(r'\s*,?\s*(?:and\s+|through\s+|to\s+|[-\u2013\u2014]\s*)?((?:19|20)\d{2})\b(?!\s*[-/\d])', re.I)


def parse_period(text: str):
    """Return (label, period_end ISO, twelve_month flag) from explicit wording only, else (None, None, False).
    Handles 'June 30, 2025 and 2026', 'September 30, 2021 through 2025' and year lists; period_end is the latest date named."""
    if not text:
        return None, None, False
    found, label_end, extra_years = [], 0, 0
    for m in DATE_RE.finditer(text):
        years, pos = [int(m.group(3))], m.end()
        while True:
            c = YEAR_CONT_RE.match(text, pos)
            if not c:
                break
            years.append(int(c.group(1))); pos = c.end()
        extra_years += len(years) - 1
        for y in years:
            try:
                found.append(date(y, MONTHS[m.group(1).capitalize()], int(m.group(2))))
            except ValueError:
                pass
        label_end = max(label_end, pos)
    if not found:
        return None, None, False
    first = next(DATE_RE.finditer(text))
    start = first.start()
    leads = list(LEAD_RE.finditer(text[:first.start()]))
    if leads and first.start() - leads[-1].start() <= 60:
        start = leads[-1].start()
    label = text_clean(text[start:label_end])
    label = label[0].upper() + label[1:] if label else label
    twelve = bool(re.search(r'(12|twelve)-month period\b', label, re.I)) and len(found) == 1
    return label, max(found).isoformat(), twelve


def twelve_month_start(period_end: str) -> str:
    d = date.fromisoformat(period_end)
    nxt = date(d.year - 1, d.month, d.day) if not (d.month == 2 and d.day == 29) else date(d.year - 1, 2, 28)
    return date.fromordinal(nxt.toordinal() + 1).isoformat()


def is_number(v):
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def xlsx_preview(path: Path):
    import io
    wb = openpyxl.load_workbook(io.BytesIO(path.read_bytes()), read_only=True, data_only=True)  # bytes: one saved file has no extension
    names = list(wb.sheetnames)
    previews, title_text = [], None
    for ws in wb.worksheets[:MAX_SHEETS]:
        try:
            ws.reset_dimensions()
        except Exception:
            pass
        rows, wide = [], False
        for raw in ws.iter_rows(values_only=True):
            if len(raw) > MAX_COLS and any(v is not None for v in raw[MAX_COLS:]):
                wide = True
            vals = [cell(v) for v in raw[:MAX_COLS]]
            while vals and vals[-1] is None:
                vals.pop()
            if vals:
                rows.append(vals)
        titles, headers, data, mode = [], [], [], 'title'
        for vals in rows:
            filled = [v for v in vals if v is not None]
            nums = [v for v in filled if is_number(v)]
            if mode == 'title':
                if len(filled) == 1 and isinstance(filled[0], str) and len(titles) < 4 and (not titles or len(filled[0]) > 40 or DATE_RE.search(filled[0]) or re.search(r'(?:19|20)\d{2}', filled[0])):
                    titles.append(filled[0])
                    continue
                mode = 'header'
            if mode == 'header':
                years_only = bool(nums) and all(isinstance(n, int) and 1900 <= n <= 2100 for n in nums) and not (isinstance(vals[0], int))
                if (not nums or years_only) and len(headers) < MAX_HEADER_ROWS:
                    headers.append(vals)
                    continue
                mode = 'data'
            if len(data) < MAX_DATA_ROWS:
                data.append(vals)
        n_data = len(rows) - len(titles) - len(headers)
        previews.append({'sheet': ws.title, 'title_lines': titles, 'header_rows': headers, 'data_rows': data,
                         'non_empty_rows_in_sheet': len(rows), 'data_rows_in_sheet': n_data,
                         'truncated': n_data > len(data) or wide, 'columns_truncated': wide})
        if title_text is None:
            title_text = ' '.join(titles) if titles else ' '.join(str(v) for r in rows[:6] for v in r if isinstance(v, str))
    wb.close()
    return names, previews, title_text


def pdf_first_page(path: Path):
    reader = PdfReader(str(path))
    try:
        text = reader.pages[0].extract_text() or ''
    except Exception:
        text = ''
    return len(reader.pages), text


# ----------------------------------------------------------------------------- CJRA per-judge tables
TOTAL_RE = re.compile(r'^Total (All Cases|Motions|Bench Trials) for ((?:[A-Za-z.]+ ){0,5}Judge)\s*:?\s+(\S.*?)\s+([\d,]+)$')
HEADER_RE = re.compile(r'^((?:District|Magistrate|Circuit|Bankruptcy|Federal Circuit|Court of International Trade|Senior District|Senior|Chief|Chief District) Judge)\s+(\S.*?)\s*$')
COURT_RE = re.compile(r'^U\.S\. (?:District Court for|Court of) .+$')
CIRCUIT_RE = re.compile(r'^(DC|\d{1,2}(?:st|nd|rd|th)|Federal) Circuit$')
DOCKET_RE = re.compile(r'^\S{1,3} \d{2}-[a-z]{2}-\d{3,6}\b')
FOOTER_RE = re.compile(r'Page:? ([\d,]+) of ([\d,]+)')


def norm_name(name: str) -> str:
    return re.sub(r'\s+', ' ', name).strip().upper()


def cjra_page_lines(path: Path, sha: str):
    """pypdf text of every page reduced to the few line kinds the parser needs; cached by file hash (extraction takes minutes)."""
    CACHE.mkdir(parents=True, exist_ok=True)
    cache = CACHE / ('%s_pages_%s.json' % (sha[:24], PARSER_VERSION))
    if cache.is_file():
        return json.loads(cache.read_text(encoding='utf-8'))
    pages = []
    for page in PdfReader(str(path)).pages:
        lines = [l.strip() for l in (page.extract_text() or '').splitlines() if l.strip()]
        pages.append({'courts': [l for l in lines if COURT_RE.match(l)], 'circuits': [l for l in lines if CIRCUIT_RE.match(l)],
                      'headers': [l for l in lines if HEADER_RE.match(l)], 'totals': [l for l in lines if l.startswith('Total ') and TOTAL_RE.match(l)],
                      'as_of': [l for l in lines if l.startswith('As of ')], 'footer': [l for l in lines if FOOTER_RE.search(l)][:1]})
    cache.write_text(json.dumps(pages), encoding='utf-8')
    return pages


def parse_cjra(path: Path, sha: str, kind: str):
    """One row per printed 'Total <kind> for <judge type> [:] <name> <count>' line. The judge header printed on the same page
    (or, on a totals-only continuation page, the latest earlier header) must name the same judge; court and circuit come from
    the page that carries the total line."""
    pages = cjra_page_lines(path, sha)
    rows, problems, asof, footer_total = [], [], Counter(), set()
    last_header, block_first_page = None, None
    for pno, pg in enumerate(pages, start=1):
        if len(pg['courts']) != 1 or len(pg['circuits']) != 1 or len(pg['headers']) > 1 or len(pg['totals']) > 1:
            problems.append({'page': pno, 'problem': 'expected one court, one circuit, at most one judge header and one total line',
                             'courts': pg['courts'][:3], 'circuits': pg['circuits'][:3], 'headers': pg['headers'][:3], 'totals': pg['totals'][:3]})
        asof.update(pg['as_of'])
        for l in pg['footer']:
            footer_total.add(int(FOOTER_RE.search(l).group(2).replace(',', '')))
        if pg['headers']:
            last_header = pg['headers'][0]
        if block_first_page is None:
            block_first_page = pno
        for l in pg['totals']:
            t = TOTAL_RE.match(l)
            if t.group(1) != kind:
                problems.append({'page': pno, 'problem': 'total line of another kind', 'line': l[:120]})
                continue
            h = HEADER_RE.match(last_header) if last_header else None
            header_name = re.sub(r'\s*\([^)]*\)\s*$', '', h.group(2)) if h else None
            ok = bool(h) and h.group(1) == t.group(2) and norm_name(t.group(3)) in (norm_name(header_name), norm_name(h.group(2)))
            rows.append({'court_as_printed': pg['courts'][0] if len(pg['courts']) == 1 else None,
                         'circuit_as_printed': pg['circuits'][0] if len(pg['circuits']) == 1 else None,
                         'judge_type_as_printed': t.group(2), 'judge_name_as_printed': text_clean(t.group(3)),
                         'judge_header_as_printed': text_clean(last_header) if last_header else None,
                         'count': int(t.group(4).replace(',', '')), 'first_page': block_first_page, 'last_page': pno,
                         'header_matches_total_line': ok})
            block_first_page = None
    return {'pages': len(pages), 'rows': rows, 'problems': problems[:50], 'problem_count': len(problems),
            'as_of_lines': dict(asof), 'footer_page_totals': sorted(footer_total), 'dangling_header': block_first_page is not None}


def parse_cjra_table2(path: Path):
    text = PdfReader(str(path)).pages[0].extract_text() or ''
    out = {}
    names = '|'.join(['Total'] + list(CIRCUITS.values()))
    for line in text.splitlines():
        m = re.match(r'^(%s) ([\d,]+) [\d.]+ [\d.]+ ([\d,]+) [\d.]+ [\d.]+ ([\d,]+) ([\d,]+) ([\d,]+)\s*$' % names, line.strip())
        if m:
            out[m.group(1)] = {'civil': int(m.group(2).replace(',', '')), 'motions': int(m.group(3).replace(',', '')), 'bench': int(m.group(4).replace(',', ''))}
    return out


# ----------------------------------------------------------------------------- main
def main():
    started = time.time()
    seeds = read_jsonl(HERE / 'seeds.jsonl')
    receipts, errors_by_seed = {}, defaultdict(list)
    for packet in ('A', 'B'):
        for r in read_jsonl(HERE / 'receipts' / ('packet_%s.jsonl' % packet)):
            if r.get('saved_path') and r.get('sha256') and r.get('seed_id') != 'robots':
                receipts[r['seed_id']] = r
            elif r.get('error'):
                errors_by_seed[r.get('seed_id')].append({'url': r.get('url'), 'error': r.get('error'), 'status': r.get('status')})
    checks, gaps, tables, ids, file_ids = [], [], [], set(), set()
    verified = 0
    for seed in seeds:
        r = receipts.get(seed['seed_id'])
        if not r:
            gaps.append({'seed_id': seed['seed_id'], 'url': seed['url'], 'listed_title': seed.get('listed_title'),
                         'reason': 'no saved file in receipts', 'receipt_errors': errors_by_seed.get(seed['seed_id'], [])})
            continue
        path = (HERE / r['saved_path']).resolve()
        if not path.is_relative_to(RAW) or not path.is_file():
            raise SystemExit('raw file missing or outside raw/: %s' % r['saved_path'])
        digest = sha256_file(path)
        if digest != r['sha256'] or path.stat().st_size != r['bytes']:
            raise SystemExit('receipt mismatch for %s' % r['saved_path'])
        if r.get('content_problem'):
            gaps.append({'seed_id': seed['seed_id'], 'url': seed['url'], 'reason': 'content problem: %s' % r['content_problem']})
            continue
        verified += 1
        fmt = r.get('detected_format') or seed.get('listed_format')
        code, series = SERIES[seed['listed_publication']]
        number = seed.get('listed_table_id')
        if number and not re.match(r'^(?:[A-Z]{1,3}-?\d+[A-Z]?|[A-Z]|CJRA [\w ]+|\d+\.\d+)$', number.strip()):
            number = None  # listing wording such as "FCMS (no table number)" is not a table number
        sheet_names, preview, page_count, file_text = [], [], None, None
        if fmt == 'xlsx':
            sheet_names, preview, file_text = xlsx_preview(path)
            file_basis = 'title line inside the workbook (first sheet)'
        elif fmt == 'pdf':
            page_count, first = pdf_first_page(path)
            file_text = ' '.join(l.strip() for l in first.splitlines() if re.match(r'\s*As of ', l)) if code == 'cjra' else None
            file_basis = '"As of" line printed on page 1 of the PDF'
        elif fmt == 'html':
            m = re.search(r'\bas of (\d{2})/(\d{2})/(\d{4})', path.read_text(encoding='utf-8', errors='replace'))
            file_text, file_basis = None, 'as-of statement printed on the saved page'
            html_asof = (m.group(0), '%s-%s-%s' % (m.group(3), m.group(1), m.group(2))) if m else None
        label, end, twelve = parse_period(file_text or '')
        basis = file_basis if end else None
        source_as_of = source_as_of_basis = None
        if fmt == 'html' and html_asof:
            label, end, basis = html_asof[0][0].upper() + html_asof[0][1:], html_asof[1], file_basis
            source_as_of, source_as_of_basis = end, file_basis
        if not end:
            label, end, twelve = parse_period(seed.get('listed_period_label') or '')
            if end:
                label = seed['listed_period_label']
                basis = 'period as listed on the publisher data-table index (seed); not stated in the file title'
                twelve = False
        if code == 'cjra' and end:
            source_as_of, source_as_of_basis = end, basis
        title_in_file = None
        if preview and preview[0]['title_lines']:
            title_in_file = ' '.join(preview[0]['title_lines'])[:600]
        slug = re.sub(r'[^A-Za-z0-9.]+', '-', number).strip('-') if number else seed['seed_id']
        table_id = '%s:%s:%s:%s' % (code, slug, end or 'undated', fmt)
        file_id = '%s-%s' % (seed['seed_id'].lower(), digest[:12])
        if table_id in ids or file_id in file_ids:
            raise SystemExit('duplicate id %s / %s' % (table_id, file_id))
        ids.add(table_id); file_ids.add(file_id)
        base = path.name.split('__', 1)[-1]
        if not base.lower().endswith('.' + fmt):
            base = '%s.%s' % (base, fmt)
        tables.append({
            'table_id': table_id, 'record_type': 'court_statistic_table', 'seed_id': seed['seed_id'], 'table_number': number,
            'title': seed['listed_title'], 'title_basis': 'title as listed on the publisher page (seed)', 'title_in_file': title_in_file,
            'series': series, 'series_code': code, 'publication_as_listed': seed['listed_publication'],
            'topic': TOPICS.get(seed.get('category'), seed.get('category')),
            'period_label': label if end else None, 'period_start': twelve_month_start(end) if (end and twelve) else None, 'period_end': end,
            'period_basis': basis, 'period_start_basis': 'computed from the explicit "12-Month Period Ending" wording' if (end and twelve) else None,
            'format': fmt, 'bytes': r['bytes'], 'sha256': digest, 'file_id': file_id, 'download_name': '%s_%s' % (seed['seed_id'], base),
            'source_page': seed.get('source_page'), 'source_url': seed['url'], 'final_url': r.get('final_url') or r.get('url'),
            'sheet_names': sheet_names, 'page_count': page_count, 'preview': preview, 'raw_path': r['saved_path'],
            'figure_label': 'publisher-reported counts%s; never a rate computed by this archive' % ((' for ' + label) if end else ' (period not stated in the file or listing)'),
            'captured_at': r.get('captured_at') or r.get('requested_at'), 'captured_at_basis': 'HTTP request timestamp in the fetch receipt',
            'source_as_of': source_as_of, 'source_as_of_basis': source_as_of_basis, 'published_at': None, 'published_at_basis': None,
            'effective_from': None, 'effective_from_basis': None, 'effective_to': None, 'effective_to_basis': None,
        })
    checks.append({'check': 'every registered raw file matches its receipt (sha256 and byte length)', 'passed': True, 'files': verified})
    checks.append({'check': 'table_id and file_id unique', 'passed': True, 'tables': len(tables)})
    xlsx = [t for t in tables if t['format'] == 'xlsx']
    checks.append({'check': 'every XLSX opened with openpyxl and produced a bounded preview', 'passed': all(t['preview'] for t in xlsx), 'xlsx': len(xlsx)})
    checks.append({'check': 'preview bounds respected', 'passed': all(len(t['preview']) <= MAX_SHEETS and all(len(p['data_rows']) <= MAX_DATA_ROWS for p in t['preview']) for t in tables)})

    # CJRA per-judge rows
    by_number = {t['table_number']: t for t in tables if t['series_code'] == 'cjra'}
    cjra_rows, cjra_stats = [], {}
    table2 = parse_cjra_table2((HERE / by_number['CJRA 2']['raw_path'])) if 'CJRA 2' in by_number else {}
    for number, (prefix, kind, count_label, t2key) in CJRA_PER_JUDGE.items():
        t = by_number.get(number)
        if not t:
            continue
        parsed = parse_cjra(HERE / t['raw_path'], t['sha256'], kind)
        rows = parsed['rows']
        sums = Counter()
        for i, row in enumerate(rows, start=1):
            circ = (row['circuit_as_printed'] or '').replace(' Circuit', '')
            sums[CIRCUITS.get(circ, circ or 'unknown')] += row['count']
            cjra_rows.append({
                'row_id': '%s:%05d' % (prefix, i), 'table_id': t['table_id'], 'table_number': number, 'file_id': t['file_id'],
                'circuit_as_printed': row['circuit_as_printed'], 'court_as_printed': row['court_as_printed'],
                'judge_type_as_printed': row['judge_type_as_printed'], 'judge_name_as_printed': row['judge_name_as_printed'],
                'count_label': count_label, 'count': row['count'], 'count_basis': 'printed "Total ... for <judge>" line in the PDF',
                'pdf_pages': [row['first_page'], row['last_page']], 'judge_header_as_printed': row['judge_header_as_printed'],
                'header_matches_total_line': row['header_matches_total_line'],
                'judge_match': {'status': 'not_attempted', 'reason': 'names are kept as printed; people are never merged by name alone'},
                'period_label': t['period_label'], 'period_start': None, 'period_end': t['period_end'],
                'captured_at': t['captured_at'], 'captured_at_basis': t['captured_at_basis'], 'source_as_of': t['source_as_of'],
                'source_as_of_basis': t['source_as_of_basis'], 'published_at': None, 'published_at_basis': None,
                'effective_from': None, 'effective_from_basis': None, 'effective_to': None, 'effective_to_basis': None})
        recon = []
        for name in ['Total'] + list(CIRCUITS.values()):
            printed = (table2.get(name) or {}).get(t2key)
            parsed_sum = sum(sums.values()) if name == 'Total' else sums.get(name, 0)
            recon.append({'circuit': name, 'sum_of_per_judge_totals': parsed_sum, 'cjra_table_2_printed': printed,
                          'match': printed is not None and printed == parsed_sum})
        header_ok = sum(1 for r in rows if r['header_matches_total_line'])
        structural = (parsed['problem_count'] == 0 and not parsed['dangling_header'] and header_ok == len(rows)
                      and parsed['footer_page_totals'] == [parsed['pages']] and len(parsed['as_of_lines']) == 1)
        cjra_stats[number] = {'pdf_pages': parsed['pages'], 'judge_rows': len(rows), 'sum_of_counts': sum(sums.values()),
                              'header_matches_total_line': header_ok, 'page_problems': parsed['problem_count'],
                              'as_of_lines': parsed['as_of_lines'], 'structural_ok': structural,
                              'reconciliation_vs_cjra_table_2': recon, 'reconciled_circuits': sum(1 for x in recon if x['match']),
                              'circuits_compared': len(recon)}
        t['cjra_per_judge'] = {k: v for k, v in cjra_stats[number].items()}
        checks.append({'check': '%s structure: one court and circuit per page, every total line preceded by the same judge header, page footer equals page count, single as-of label' % number,
                       'passed': structural, 'judge_rows': len(rows), 'pdf_pages': parsed['pages']})
        checks.append({'check': '%s reconciliation: per-circuit sum of printed per-judge totals equals CJRA Table 2 (%s)' % (number, t2key),
                       'passed': all(x['match'] for x in recon), 'matched': sum(1 for x in recon if x['match']), 'compared': len(recon)})

    def dump(name, rows):
        payload = ''.join(json.dumps(r, ensure_ascii=False, sort_keys=True) + '\n' for r in rows).encode('utf-8')
        (HERE / name).write_bytes(payload)
        return {'path': name, 'sha256': hashlib.sha256(payload).hexdigest(), 'rows': len(rows)}

    data_files = [dump('tables.jsonl', tables), dump('cjra_rows.jsonl', cjra_rows)]
    (HERE / 'gaps.json').write_text(json.dumps({'not_saved': gaps}, indent=1), encoding='utf-8')
    counts = {
        'tables': len(tables), 'by_series': dict(Counter(t['series'] for t in tables)), 'by_format': dict(Counter(t['format'] for t in tables)),
        'by_topic': dict(Counter(t['topic'] for t in tables)), 'with_period_end': sum(1 for t in tables if t['period_end']),
        'period_from_file': sum(1 for t in tables if t['period_basis'] and 'seed' not in t['period_basis']),
        'period_from_listing': sum(1 for t in tables if t['period_basis'] and 'seed' in t['period_basis']),
        'period_unknown': sum(1 for t in tables if not t['period_end']), 'xlsx_previews': len(xlsx),
        'preview_sheets': sum(len(t['preview']) for t in tables), 'cjra_rows': len(cjra_rows),
        'cjra_rows_by_table': dict(Counter(r['table_number'] for r in cjra_rows)), 'seeds': len(seeds), 'seeds_not_saved': len(gaps),
        'cjra': cjra_stats,
    }
    hard = [c for c in checks if 'reconciliation' not in c['check']]
    passed = all(c['passed'] for c in hard) and len(tables) > 0
    validation = {
        'schema_version': '1', 'status': 'passed' if passed else 'failed', 'ready': bool(passed),
        'validated_at': datetime.now(timezone.utc).strftime('%Y-%m-%dT%H:%M:%SZ'), 'data_files': data_files, 'counts': counts, 'checks': checks,
        'qualification': ('Registered originals of Administrative Office of the U.S. Courts statistical tables saved from uscourts.gov on 2026-09-19 '
                          '(%d files; latest periods only, no historical backfill). Every figure is a publisher-reported count for the stated period or '
                          'as-of date, never a rate computed here; CJRA per-judge rows keep judge names as printed and are not linked to judge profiles.' % len(tables)),
        'license_ref': 'Works of the U.S. federal judiciary (Administrative Office of the U.S. Courts), published at uscourts.gov; no copyright notice on the tables. Not independently reviewed.',
        'inputs': [{'path': 'seeds.jsonl', 'sha256': sha256_file(HERE / 'seeds.jsonl')},
                   {'path': 'receipts/packet_A.jsonl', 'sha256': sha256_file(HERE / 'receipts/packet_A.jsonl')},
                   {'path': 'receipts/packet_B.jsonl', 'sha256': sha256_file(HERE / 'receipts/packet_B.jsonl')}],
    }
    (HERE / 'validation.json').write_text(json.dumps(validation, indent=1, ensure_ascii=False), encoding='utf-8')
    write_spec()
    print(json.dumps({'status': validation['status'], 'tables': len(tables), 'cjra_rows': len(cjra_rows), 'gaps': len(gaps),
                      'period_unknown': counts['period_unknown'], 'seconds': round(time.time() - started, 1)}, indent=1))
    for c in checks:
        print(('PASS ' if c['passed'] else 'FAIL ') + c['check'], {k: v for k, v in c.items() if k not in ('check', 'passed')})


def write_spec():
    example_listing = example_detail = None
    try:
        sys.path.insert(0, str(HERE.parents[1] / 'delivery' / 'archive-directory'))
        import court_statistics as adapter  # noqa
        example_listing = adapter.listing({'series': 'stfj', 'limit': '2'})
        first = (example_listing.get('results') or [{}])[0].get('id')
        example_detail = adapter.detail(first) if first else None
        if example_detail:
            for s in example_detail.get('sections', []):
                if 'rows' in s:
                    s['rows'] = s['rows'][:3]
    except Exception as exc:  # adapter not written yet or gate closed
        example_listing = {'note': 'adapter not importable at build time: %s' % exc.__class__.__name__}
    spec = {
        'adapter': 'delivery/archive-directory/court_statistics.py', 'area_name': 'court_statistics',
        'functions': {'listing': 'listing(params: dict) -> dict  # q, series, topic, format, period, page, limit<=100',
                      'detail': 'detail(id: str) -> dict | None  # id = table_id, or "<table_id>@<court slug>" for the CJRA per-judge rows of one court',
                      'original': 'original(file_id: str) -> (bytes, mime, filename) | None  # re-hashed at serve time, confined to raw/'},
        'routes': [{'method': 'GET', 'path': '/api/area/court_statistics', 'params': ['q', 'series', 'topic', 'format', 'period', 'page', 'limit']},
                   {'method': 'GET', 'path': '/api/area/court_statistics/item', 'params': ['id']},
                   {'method': 'GET', 'path': '/supplement-files/court_statistics/<file_id>'}],
        'ui': {'area': 'Court statistics', 'columns': ['Table', 'Series', 'Period', 'Format'],
               'filters': ['Search', 'Series', 'Topic', 'Format', 'Period'],
               'caveats': ['Publisher-reported counts for the stated period or as-of date; this archive computes no rates.',
                           'Percent columns inside a preview are the publisher\'s own cells, shown as printed.',
                           'Civil counts include MDL member cases, which can dominate a district\'s totals.',
                           'CJRA judge names are as printed; they are not linked to judge profiles.']},
        'edges': 'none emitted (no evidence-based joins were attempted in this round)',
        'example_listing': example_listing, 'example_detail': example_detail,
    }
    (HERE / 'integration_spec.json').write_text(json.dumps(spec, indent=1, ensure_ascii=False), encoding='utf-8')


if __name__ == '__main__':
    main()
